from .base import BaseParser, ParsedDocument


class DocxParser(BaseParser):
    def parse(self, file_path: str, document_id: str, metadata: dict | None = None) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx is required for DOCX parsing. Install with: pip install python-docx")

        doc = Document(file_path)
        full_text_parts: list[str] = []
        sections: list[dict] = []
        tables: list[dict] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Detect heading by style
            if para.style.name.startswith("Heading"):
                try:
                    level = int(para.style.name.split()[-1])
                except ValueError:
                    level = 1
                sections.append({
                    "level": level,
                    "title": text,
                    "section_path": text,
                })

            full_text_parts.append(text)

        # Extract tables
        for i, table in enumerate(doc.tables):
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                tables.append({
                    "index": i,
                    "markdown": "\n".join(rows),
                })

        full_text = "\n".join(full_text_parts)

        return ParsedDocument(
            document_id=document_id,
            file_path=file_path,
            title=metadata.get("title", ""),
            sections=sections,
            tables=tables,
            full_text=full_text,
            metadata=metadata or {},
        )
